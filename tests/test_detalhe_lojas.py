"""Regressao da contagem de lojas no detalhe e no Excel (sem banco)."""
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from services.detalhe_service import obter_lojas
from services.exportar_arvore import _com_movimento, agrupar, tem_movimento
from services.exportar_service import exportar_detalhe


def loja(chave, **extra):
    return {
        "chave_loja": chave, "gerencia": "Gestão", "coordenacao": "III", "supervisao": "Comercial",
        "ativa": True, "qtd_tentativas": 0, "qtd_convertidas": 0, "pct_conversao": 0,
        "qtd_averbado": 0, "vlr_averbado": 0, "qtd_aguardando": 0, "vlr_aguardando": 0,
        "qtd_nao_averbado": 0, "vlr_nao_averbado": 0, **extra,
    }


class TestLojasNoRecorte(TestCase):
    def test_multiplas_operacoes_inativa_e_so_tentativas(self):
        lojas = [loja(1, qtd_averbado=3), loja(2, ativa=False, qtd_averbado=2), loja(3, qtd_tentativas=10)]
        grupo = agrupar(lojas, "gerencia")[0]
        self.assertEqual(grupo["qtd_averbado"], 5)
        self.assertEqual(grupo["qtd_lojas_mov"], 2)

    def test_chave_repetida_conta_uma_loja(self):
        repetida = loja(1, qtd_averbado=2)
        self.assertEqual(agrupar([repetida, repetida], "gerencia")[0]["qtd_lojas_mov"], 1)

    def test_situacoes_e_tentativas(self):
        lojas = [loja(1, qtd_averbado=2), loja(2, qtd_aguardando=1), loja(3, qtd_nao_averbado=1), loja(4, qtd_tentativas=5, ativa=False)]
        for situacao, chaves in (("AVERBADO", [1]), ("AGUARDANDO AVERBACAO", [2]), ("NAO AVERBADO", [3]), ("PENDENTE", [2, 3])):
            with self.subTest(situacao=situacao):
                self.assertEqual([l["chave_loja"] for l in lojas if _com_movimento(l, situacao=situacao)], chaves)
                self.assertEqual(agrupar(lojas, "gerencia", situacao=situacao)[0]["qtd_lojas_mov"], len(chaves))
        self.assertEqual([l["chave_loja"] for l in lojas if tem_movimento(l)], [1])
        self.assertEqual(agrupar(lojas, "gerencia")[0]["qtd_lojas_mov"], 1)

    def test_tentativas_contam_lojas_que_tentaram_inclusive_inativas(self):
        lojas = [loja(1, qtd_averbado=2), loja(2, qtd_tentativas=8), loja(3, qtd_tentativas=1, ativa=False)]
        grupo = agrupar(lojas, "gerencia", foco="tentativas")[0]
        self.assertEqual(grupo["qtd_lojas_mov"], 2)

    @patch("services.detalhe_service.run_query", return_value=[])
    def test_api_filtra_por_producao_sem_cortar_ativas(self, query):
        obter_lojas({"produto": "INSS", "situacao": "AVERBADO", "data_ini": "2026-09-24", "data_fim": "2026-09-24"})
        _, params, tokens = query.call_args.args
        self.assertIn("CRÉDITO CONSIGNADO INSS", params)
        self.assertIn("AVERBADO", params)
        self.assertIn("P.QTD_AVERBADO", tokens["SOMENTE_COM_MOVIMENTO"])
        self.assertNotIn("T.", tokens["SOMENTE_COM_MOVIMENTO"])
        self.assertNotIn("ATIVA", tokens["SOMENTE_COM_MOVIMENTO"])

    @patch("services.detalhe_service.run_query", return_value=[])
    def test_incluir_sem_movimento_preserva_universo_completo(self, query):
        obter_lojas({"situacao": "AVERBADO", "incluir_sem_movimento": "1"})
        self.assertEqual(query.call_args.args[2]["SOMENTE_COM_MOVIMENTO"], "")

    @patch("services.detalhe_service.run_query", return_value=[])
    def test_lista_geral_sem_situacao_usa_lojas_averbadas(self, query):
        obter_lojas({"produto": "PRIVADO"})
        _, params, tokens = query.call_args.args
        self.assertIn("CRÉDITO CONSIGNADO PRIVADO", params)
        self.assertEqual(tokens["SOMENTE_COM_MOVIMENTO"], "WHERE P.QTD_AVERBADO > 0")

    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_arvore.obter_lojas")
    def test_excel_resumo_e_lista_de_lojas_usam_mesmo_recorte(self, obter, _query):
        obter.return_value = [loja(1, qtd_averbado=3), loja(2, ativa=False, qtd_averbado=2), loja(3, qtd_tentativas=10)]
        _, conteudo = exportar_detalhe({"produto": "INSS", "situacao": "AVERBADO"})
        wb = load_workbook(BytesIO(conteudo))
        self.assertEqual(wb["Ger. Gestão"]["B2"].value, 2)
        self.assertEqual([row[3].value for row in wb["Lojas"].iter_rows(min_row=2)], [1, 2])

    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_arvore.obter_lojas")
    def test_excel_movimento_e_lista_acompanham_assunto_da_tela(self, obter, _query):
        obter.return_value = [loja(1, qtd_averbado=3), loja(2, qtd_tentativas=10, qtd_aguardando=2), loja(3, qtd_nao_averbado=1, ativa=False, gerencia="Outra gestão")]
        for args, chaves in (({}, [1]), ({"foco": "tentativas"}, [2]), ({"situacao": "AGUARDANDO AVERBACAO"}, [2]), ({"situacao": "NAO AVERBADO"}, [3]), ({"situacao": "PENDENTE", "produto": "INSS"}, [2, 3])):
            with self.subTest(args=args):
                _, conteudo = exportar_detalhe(args)
                wb = load_workbook(BytesIO(conteudo))
                self.assertEqual(sum(row[1].value for row in wb["Ger. Gestão"].iter_rows(min_row=2)), len(chaves))
                self.assertEqual([row[3].value for row in wb["Lojas"].iter_rows(min_row=2)], chaves)
