"""Testes do Excel por etapas (sem banco).

Rodar: python -m unittest tests.test_exportar_xlsx
"""
from datetime import date
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from services.exportar_arvore import agrupar, nivel_inicial, niveis_a_partir, tem_movimento
from services.exportar_service import exportar_detalhe, exportar_du, exportar_equipe, exportar_tabela
from services.xlsx_planilha import Coluna, VERMELHO, bytes_workbook, escrever_planilha, novo_workbook


def _abrir(conteudo):
    return load_workbook(BytesIO(conteudo))


def _produtos():
    return {
        p: {
            "vlr": 10.0, "qtd": 1, "lojas": 1, "vlr_aguardando": 0.0,
            "tentativas": 2, "convertidas": 1, "pct_conversao": 50.0,
        }
        for p in ("INSS", "PRIVADO", "PUBLICO")
    }


def _equipe(nivel, rotulo):
    return {
        "nivel": nivel,
        "rotulo": rotulo,
        "linhas": [{
            "descricao": "Sul",
            "pai_gerencia": "Gestão Sul",
            "pai_coordenacao": "III Sul",
            "qtd_lojas": 2,
            "qtd_lojas_producao": 1,
            "pct_cobertura": 50.0,
            "produtos": _produtos(),
        }],
    }


def _loja(**extra):
    base = {
        "chave_loja": 1, "nome_loja": "Loja Centro", "municipio": "São Paulo", "uf": "SP",
        "gerencia": "Sul", "coordenacao": "III Sul", "supervisao": "Comercial Sul",
        "qtd_averbado": 1, "vlr_averbado": 100.0,
        "qtd_aguardando": 0, "vlr_aguardando": 0.0,
        "qtd_nao_averbado": 0, "vlr_nao_averbado": 0.0,
        "ativa": True, "qtd_tentativas": 2, "qtd_convertidas": 1, "pct_conversao": 50.0,
    }
    base.update(extra)
    return base


def _ranking(nivel, rotulo, pais):
    return {
        "nivel": nivel,
        "rotulo": rotulo,
        "pais": pais,
        "metrica": "vlr",
        "mes_ref": 202609,
        "du": 5,
        "linhas": [{
            "descricao": "Sul",
            "pai_gerencia": "Gestão Sul",
            "pai_coordenacao": "III Sul",
            "pai_supervisao": "Comercial Sul",
            "atual_du": 10.0, "media_du": 8.0,
            "atual_acum": 50.0, "media_acum": 40.0,
            "diferenca": 10.0, "desvio_pct": 25.0,
            "lojas_acum": 3, "media_lojas": 2.5,
            "ultimo_du": 5, "dus_parado": 0, "status": "acima",
        }],
    }


class TestPlanilha(TestCase):
    def test_cabecalho_vermelho_freeze_filtro_e_formato(self):
        wb = novo_workbook()
        escrever_planilha(wb, "Ger. Gestão", [
            Coluna("Nome", lambda l: l["nome"]),
            Coluna("Valor", lambda l: l["vlr"], "reais"),
        ], [{"nome": "Sul", "vlr": 10.5}])
        ws = _abrir(bytes_workbook(wb))["Ger. Gestão"]
        self.assertEqual(ws["A1"].value, "Nome")
        self.assertIn(VERMELHO, ws["A1"].fill.fgColor.rgb or "")
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.auto_filter.ref.startswith("A1"))
        self.assertEqual(ws["B2"].value, 10.5)
        self.assertEqual(ws["B2"].number_format, '"R$" #,##0.00')
        self.assertEqual(ws["B2"].alignment.horizontal, "center")
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 10)
        self.assertLessEqual(ws.column_dimensions["A"].width, 42)

    def test_qtd_no_cabecalho_fica_maiusculo(self):
        wb = novo_workbook()
        escrever_planilha(wb, "Lojas", [
            Coluna("Averbado (QTD)", lambda l: l["qtd"], "inteiro"),
        ], [{"qtd": 3}])
        ws = _abrir(bytes_workbook(wb))["Lojas"]
        self.assertEqual(ws["A1"].value, "Averbado (QTD)")
        self.assertEqual(ws["A2"].alignment.horizontal, "center")


class TestNiveis(TestCase):
    def test_nivel_inicial_segue_o_filtro(self):
        self.assertEqual(nivel_inicial({}), "gerencia")
        self.assertEqual(nivel_inicial({"gerencia": "1"}), "coordenacao")
        self.assertEqual(nivel_inicial({"coordenacao": "2"}), "supervisao")
        self.assertEqual(nivel_inicial({"supervisao": "3"}), "loja")
        self.assertEqual(nivel_inicial({"loja": "4"}), "loja")

    def test_niveis_pulam_os_ja_fixados(self):
        ordem = ("gerencia", "coordenacao", "supervisao")
        self.assertEqual(niveis_a_partir({}, ordem), ["gerencia", "coordenacao", "supervisao"])
        self.assertEqual(niveis_a_partir({"gerencia": "1"}, ordem), ["coordenacao", "supervisao"])
        self.assertEqual(niveis_a_partir({"supervisao": "3"}, ordem), [])


class TestAgrupar(TestCase):
    def test_soma_metricas_e_mantem_pais(self):
        lojas = [
            _loja(coordenacao="III A", supervisao="C1", vlr_averbado=100, qtd_averbado=1, qtd_tentativas=4, qtd_convertidas=2),
            _loja(chave_loja=2, coordenacao="III A", supervisao="C2", vlr_averbado=50, qtd_averbado=1, qtd_tentativas=0, qtd_convertidas=0),
        ]
        grupos = agrupar(lojas, "coordenacao")
        self.assertEqual(len(grupos), 1)
        self.assertEqual(grupos[0]["gerencia"], "Sul")
        self.assertEqual(grupos[0]["coordenacao"], "III A")
        self.assertEqual(grupos[0]["vlr_averbado"], 150)
        self.assertEqual(grupos[0]["qtd_lojas_mov"], 2)
        self.assertEqual(grupos[0]["pct_conversao"], 50.0)

    def test_loja_sem_movimento(self):
        self.assertFalse(tem_movimento(_loja(
            qtd_averbado=0, qtd_aguardando=0, qtd_nao_averbado=0, qtd_tentativas=0,
        )))


class TestComposicao(TestCase):
    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_arvore.obter_lojas", return_value=[_loja()])
    @patch("services.exportar_service.obter_equipe")
    def test_equipe_uma_aba_por_etapa(self, mock_equipe, _lojas, _query):
        mock_equipe.side_effect = [
            _equipe("gerencia", "Ger. Gestão"),
            _equipe("coordenacao", "Ger. Comercial III"),
            _equipe("supervisao", "Ger. Comercial"),
        ]
        nome, conteudo = exportar_equipe({"data_ini": "2026-09-01", "data_fim": "2026-09-25"})
        wb = _abrir(conteudo)
        self.assertEqual(nome, "equipe_20260901_20260925.xlsx")
        self.assertEqual(wb.sheetnames, [
            "Ger. Gestão", "Ger. Comercial III", "Ger. Comercial",
            "Lojas", "Contratos", "Tentativas",
        ])
        self.assertEqual(wb["Ger. Comercial III"]["A1"].value, "Ger. Gestão")
        self.assertEqual(wb["Lojas"]["A1"].value, "Ger. Gestão")
        self.assertEqual(wb["Contratos"]["A1"].value, "Ger. Gestão")
        self.assertIn(VERMELHO, wb["Ger. Gestão"]["A1"].fill.fgColor.rgb or "")
        cab_lojas = [c.value for c in wb["Lojas"][1]]
        self.assertIn("Averbado (QTD)", cab_lojas)
        self.assertEqual(wb["Lojas"]["K2"].number_format, '"R$" #,##0.00')
        self.assertEqual(wb["Lojas"]["K2"].alignment.horizontal, "center")

    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_arvore.obter_lojas", return_value=[_loja()])
    @patch("services.exportar_service.obter_equipe")
    def test_equipe_pula_nivel_ja_filtrado(self, mock_equipe, _lojas, _query):
        mock_equipe.side_effect = [
            _equipe("coordenacao", "Ger. Comercial III"),
            _equipe("supervisao", "Ger. Comercial"),
        ]
        _nome, conteudo = exportar_equipe({
            "gerencia": "10", "data_ini": "2026-09-01", "data_fim": "2026-09-25",
        })
        self.assertEqual(_abrir(conteudo).sheetnames, [
            "Ger. Comercial III", "Ger. Comercial", "Lojas", "Contratos", "Tentativas",
        ])

    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_arvore.obter_lojas", return_value=[_loja()])
    def test_detalhe_arvore_completa(self, _lojas, _query):
        _nome, conteudo = exportar_detalhe({"data_ini": "2026-09-01", "data_fim": "2026-09-25"})
        self.assertEqual(_abrir(conteudo).sheetnames, [
            "Ger. Gestão", "Ger. Comercial III", "Ger. Comercial",
            "Lojas", "Contratos", "Tentativas",
        ])

    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_arvore.obter_lojas", return_value=[_loja()])
    @patch("services.exportar_service.obter_tabela", return_value=[{
        "dia": "2026-09-25",
        "produtos": {p: {"qtd": 1, "lojas": 1, "vlr": 10.0, "tentativas": 2, "convertidas": 1}
                     for p in ("INSS", "PRIVADO", "PUBLICO")},
        "total_qtd": 3, "total_lojas": 1, "total_vlr": 30.0,
        "qtd_aguardando": 0, "lojas_aguardando": 0, "vlr_aguardando": 0.0,
        "qtd_nao_averbado": 0, "lojas_nao_averbado": 0, "vlr_nao_averbado": 0.0,
    }])
    def test_tabela_comeca_por_dia(self, _tabela, _lojas, _query):
        _nome, conteudo = exportar_tabela({"data_ini": "2026-09-01", "data_fim": "2026-09-25"})
        self.assertEqual(_abrir(conteudo).sheetnames[0], "Por dia")
        self.assertIn("Contratos", _abrir(conteudo).sheetnames)
        self.assertIn("Tentativas", _abrir(conteudo).sheetnames)

    @patch("services.exportar_arvore.run_query", return_value=[])
    @patch("services.exportar_service.run_query", return_value=[{"DU": 5, "DATA": date(2026, 9, 5)}])
    @patch("services.exportar_service.du_selecionado", return_value=5)
    @patch("services.exportar_service.carregar_calendario", return_value={
        "mes_ref": 202609, "meses": [202607, 202608, 202609],
    })
    @patch("services.exportar_service.obter_ranking")
    def test_du_ranking_depois_contratos(self, mock_rank, _cal, _du, _cal_q, _arv_q):
        mock_rank.side_effect = [
            _ranking("gerencia", "Ger. Gestão", []),
            _ranking("coordenacao", "Ger. Comercial III", ["gerencia"]),
            _ranking("supervisao", "Ger. Comercial", ["gerencia", "coordenacao"]),
            _ranking("loja", "Loja", ["gerencia", "coordenacao", "supervisao"]),
        ]
        nome, conteudo = exportar_du({"metrica": "vlr"})
        wb = _abrir(conteudo)
        self.assertEqual(nome, "dia_util_vlr_202609_du5.xlsx")
        self.assertEqual(wb.sheetnames, [
            "Ger. Gestão", "Ger. Comercial III", "Ger. Comercial",
            "Lojas", "Contratos", "Tentativas",
        ])
        self.assertEqual(wb["Ger. Comercial III"]["A1"].value, "Ger. Gestão")
        self.assertEqual(wb["Lojas"]["A1"].value, "Ger. Gestão")


if __name__ == "__main__":
    import unittest
    unittest.main()
