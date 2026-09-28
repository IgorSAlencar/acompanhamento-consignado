"""Escrita de planilhas Excel com cabecalho da marca e largura ajustada."""
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

VERMELHO = "CC092F"
LARGURA_MIN = 10
LARGURA_MAX = 42

_PREENCHIMENTO = PatternFill("solid", fgColor=VERMELHO)
_FONTE = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_ALINHAMENTO_CABECALHO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALINHAMENTO_NUMERO = Alignment(horizontal="center", vertical="center")

_FORMATOS = {
    "reais": '"R$" #,##0.00',
    "inteiro": "#,##0",
    "percentual": "0.0",
    "data": "DD/MM/YYYY",
    "decimal": "#,##0.00",
}


@dataclass(frozen=True)
class Coluna:
    rotulo: str
    valor: Callable
    tipo: str = "texto"


def novo_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def bytes_workbook(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def escrever_planilha(wb: Workbook, nome: str, colunas: list[Coluna], linhas: list) -> Worksheet:
    """Cria uma aba com cabecalho vermelho, freeze, autofiltro e largura pela coluna."""
    ws = wb.create_sheet(nome[:31])
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for i, coluna in enumerate(colunas, start=1):
        celula = ws.cell(1, i, coluna.rotulo)
        celula.fill = _PREENCHIMENTO
        celula.font = _FONTE
        celula.alignment = _ALINHAMENTO_CABECALHO

    larguras = [len(c.rotulo) for c in colunas]
    for linha in linhas:
        valores = [coluna.valor(linha) for coluna in colunas]
        ws.append([_celula(coluna.tipo, valor) for coluna, valor in zip(colunas, valores)])
        for i, (coluna, valor) in enumerate(zip(colunas, valores)):
            larguras[i] = max(larguras[i], len(_texto(valor)))

    for i, (coluna, largura) in enumerate(zip(colunas, larguras), start=1):
        extra = 4 if coluna.tipo == "reais" else 2
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = min(max(largura + extra, LARGURA_MIN), LARGURA_MAX)
        if coluna.tipo not in _FORMATOS:
            continue
        fmt = _FORMATOS[coluna.tipo]
        for linha_n in range(2, ws.max_row + 1):
            celula = ws.cell(linha_n, i)
            celula.number_format = fmt
            celula.alignment = _ALINHAMENTO_NUMERO

    ultima = get_column_letter(len(colunas)) if colunas else "A"
    ws.auto_filter.ref = f"A1:{ultima}{max(ws.max_row, 1)}"
    return ws


def _celula(tipo: str, valor):
    if valor is None:
        return None
    if tipo == "data":
        return _como_data(valor)
    if tipo in ("reais", "inteiro", "percentual", "decimal"):
        return valor if isinstance(valor, (int, float)) else 0
    return valor


def _como_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str) and len(valor) >= 10:
        return date.fromisoformat(valor[:10])
    return valor


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)
